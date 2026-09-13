import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from data_loader import get_trading_universe, download_historical_data
from strategy import CrossSectionalSwingEngine

def export_trades_to_excel(trades_df, filename="Actual_Trade_Log.xlsx"):
    """Generates the styled Excel trade log with conditional red/green coloring."""
    if trades_df.empty:
        print("No trades were executed to log.")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Trade Log"

    cols = list(trades_df.columns)
    ws.append(cols)

    # Add data
    for r in trades_df.values.tolist():
        ws.append(r)

    # Styles
    header_fill = PatternFill(start_color="2c3e50", end_color="2c3e50", fill_type="solid")
    header_font = Font(color="ffffff", bold=True)
    profit_fill = PatternFill(start_color="e6f4ea", end_color="e6f4ea", fill_type="solid")
    loss_fill = PatternFill(start_color="fce8e6", end_color="fce8e6", fill_type="solid")
    thin_border = Border(left=Side(style='thin', color='e0e0e0'),
                         right=Side(style='thin', color='e0e0e0'),
                         top=Side(style='thin', color='e0e0e0'),
                         bottom=Side(style='thin', color='e0e0e0'))

    # Apply Header Styles
    for col in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Apply Row Styles based on PnL
    for row in range(2, ws.max_row + 1):
        pnl_cell = ws.cell(row=row, column=cols.index('Gross PnL (₹)') + 1)
        is_profit = pnl_cell.value > 0
        
        for col in range(1, len(cols) + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border
            
            # Color code
            cell.fill = profit_fill if is_profit else loss_fill
                
            # Number formatting
            if cols[col-1] in ['Avg Entry Price', 'Exit Price', 'Gross PnL (₹)']:
                cell.number_format = '#,##0.00'
            elif cols[col-1] == 'ROI (%)':
                cell.number_format = '0.00'

    # Adjust Column Widths
    for col in range(1, len(cols) + 1):
        col_letter = get_column_letter(col)
        ws.column_dimensions[col_letter].width = 16

    ws.freeze_panes = "A2"
    wb.save(filename)
    print(f"\n=> Excel Trade Log saved to {filename}")


def calculate_metrics(equity_curve):
    """Calculates basic performance metrics from the equity curve."""
    initial = equity_curve['Portfolio_Value'].iloc[0]
    final = equity_curve['Portfolio_Value'].iloc[-1]
    total_return = ((final - initial) / initial) * 100
    
    peak = equity_curve['Portfolio_Value'].cummax()
    drawdown = (equity_curve['Portfolio_Value'] - peak) / peak
    max_drawdown = drawdown.min() * 100
    
    print("\n" + "="*30)
    print("BACKTEST RESULTS")
    print("="*30)
    print(f"Initial Capital:  ₹{initial:,.2f}")
    print(f"Final Capital:    ₹{final:,.2f}")
    print(f"Total Return:     {total_return:.2f}%")
    print(f"Max Drawdown:     {max_drawdown:.2f}%")
    print("="*30)

def main():
    universe = get_trading_universe("Filtered_Best_ETFs_For_Trading.xlsx")
    print(f"Fetching data for {len(universe)} ETFs...")
    raw_data = download_historical_data(universe, period="3y", interval="1d")
    
    data_dict = {}
    for ticker in universe:
        try:
            if isinstance(raw_data.columns, pd.MultiIndex):
                if ticker in raw_data.columns.get_level_values(0):
                    df = raw_data[ticker].copy()
                elif ticker in raw_data.columns.get_level_values(1):
                    df = raw_data.xs(ticker, level=1, axis=1).copy()
                else:
                    continue
            else:
                df = raw_data.copy()
            
            if 'Close' in df.columns:
                df = df.dropna(subset=['Close'])
                
            if not df.empty and len(df) > 50: 
                data_dict[ticker] = df
                
        except Exception as e:
            pass 
            
    print(f"Valid ETFs loaded: {len(data_dict)}")
    engine = CrossSectionalSwingEngine(data_dict, initial_capital=100000)
    
    # UNPACK BOTH DATA FRAMES HERE
    equity_curve, trades_df = engine.run_backtest()
    
    if not equity_curve.empty:
        calculate_metrics(equity_curve)
        
        # Export the real trade log to Excel
        export_trades_to_excel(trades_df, "Actual_Trade_Log.xlsx")
        
        plt.figure(figsize=(12, 6))
        plt.plot(equity_curve.index, equity_curve['Portfolio_Value'], label='Strategy Equity')
        plt.title('ETF Swing Strategy')
        plt.xlabel('Date')
        plt.ylabel('Portfolio Value')
        plt.legend()
        plt.grid(True)
        plt.savefig("equity_curve.png")
        print("=> Plot saved as equity_curve.png")
    else:
        print("Error: Equity curve is empty.")

if __name__ == "__main__":
    main()
